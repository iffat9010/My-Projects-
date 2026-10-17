import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
from datetime import datetime, date
import os

FILE = "blood_donors.xlsx"

HEADERS = [
    "Registered On", "Name", "Age", "Mobile",
    "City", "Blood Group", "Last Donation", "Available"
]

GROUPS = ["A+", "A-", "B+", "B-", "AB+", "AB-", "O+", "O-"]

RED = "red"
DARK_RED = "darkred"
BG = "white"
WHITE = "white"


# Create or open Excel file
def get_sheet():
    if os.path.exists(FILE):
        wb = load_workbook(FILE)

        if "Donors" in wb.sheetnames:
            ws = wb["Donors"]
        else:
            ws = wb.create_sheet("Donors")
            ws.append(HEADERS)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Donors"
        ws.append(HEADERS)

    return wb, ws


# Clear registration fields
def clear_form():
    name_var.set("")
    age_var.set("")
    mobile_var.set("")
    city_var.set("")
    blood_var.set("")
    donation_var.set("")
    available_var.set("Yes")


# Register a new donor
def register_donor():
    name = name_var.get().strip()
    age = age_var.get().strip()
    mobile = mobile_var.get().strip()
    city = city_var.get().strip()
    blood = blood_var.get()
    donation = donation_var.get().strip()
    available = available_var.get()

    if not all([name, age, mobile, city, blood, available]):
        messagebox.showwarning(
            "Missing Details", "Please fill in all required fields."
        )
        return

    if not age.isdigit() or not 18 <= int(age) <= 100:
        messagebox.showerror(
            "Invalid Age", "Age must be between 18 and 100."
        )
        return

    if not (
        mobile.isdigit()
        and len(mobile) == 10
        and mobile[0] in "6789"
    ):
        messagebox.showerror(
            "Invalid Mobile",
            "Enter a valid 10-digit Indian mobile number."
        )
        return

    if donation:
        try:
            donation_date = datetime.strptime(
                donation, "%d-%m-%Y"
            ).date()

            if donation_date > date.today():
                raise ValueError

        except ValueError:
            messagebox.showerror(
                "Invalid Date",
                "Enter the last donation date as DD-MM-YYYY."
            )
            return

    try:
        wb, ws = get_sheet()

        # Prevent duplicate mobile numbers
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) > 3 and str(row[3] or "").strip() == mobile:
                wb.close()
                messagebox.showwarning(
                    "Duplicate", "This mobile number is already registered."
                )
                return

        ws.append([
            datetime.now().strftime("%d-%m-%Y %H:%M"),
            name,
            int(age),
            mobile,
            city,
            blood,
            donation if donation else "Not provided",
            available
        ])

        wb.save(FILE)
        wb.close()

        messagebox.showinfo(
            "Success", "Donor registered and saved in Excel!"
        )

        clear_form()
        show_all()

    except PermissionError:
        messagebox.showerror(
            "File Error",
            "Please close the Excel file and try again."
        )

    except Exception as error:
        messagebox.showerror(
            "Error", f"Could not save the record:\n{error}"
        )


# Display donor records
def show_results(rows, message):
    for item in table.get_children():
        table.delete(item)

    for row in rows:
        table.insert(
            "", "end",
            values=(row[1], row[2], row[4], row[5], row[7])
        )

    status.config(text=message)


# Show all registered donors
def show_all():
    try:
        wb, ws = get_sheet()

        rows = [
            row for row in ws.iter_rows(min_row=2, values_only=True)
            if row and any(value is not None for value in row)
        ]

        wb.close()
        show_results(rows, f"Total registered donors: {len(rows)}")

    except Exception as error:
        messagebox.showerror(
            "Error", f"Could not read Excel:\n{error}"
        )


# Search for available donors
def search_donors():
    blood = search_blood.get()
    city = search_city.get().strip().lower()

    if blood == "All":
        blood = ""

    if not blood and not city:
        messagebox.showwarning(
            "Search Required",
            "Select a blood group or enter a city."
        )
        return

    try:
        wb, ws = get_sheet()
        matches = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(value is not None for value in row):
                continue

            donor_blood = str(row[5] or "").strip()
            donor_city = str(row[4] or "").strip().lower()
            available = str(row[7] or "").strip().lower()

            if (
                available == "yes"
                and (not blood or donor_blood == blood)
                and (not city or city in donor_city)
            ):
                matches.append(row)

        wb.close()

        show_results(
            matches, f"Available matching donors: {len(matches)}"
        )

        if not matches:
            messagebox.showinfo(
                "No Results",
                "No matching available donors were found."
            )

    except Exception as error:
        messagebox.showerror(
            "Error", f"Search failed:\n{error}"
        )


# ---------------- MAIN GUI ----------------

root = tk.Tk()
root.title("Blood Donor Registration and Emergency Search")
root.geometry("900x700")
root.minsize(820, 650)
root.configure(bg=BG)

style = ttk.Style()
style.theme_use("clam")

style.configure(
    "Treeview",
    font=("Arial", 10),
    rowheight=27,
    background=WHITE,
    fieldbackground=WHITE
)

style.configure(
    "Treeview.Heading",
    font=("Arial", 10, "bold"),
    background="pink",
    foreground=DARK_RED
)

name_var = tk.StringVar()
age_var = tk.StringVar()
mobile_var = tk.StringVar()
city_var = tk.StringVar()
blood_var = tk.StringVar()
donation_var = tk.StringVar()
available_var = tk.StringVar(value="Yes")

search_blood = tk.StringVar(value="All")
search_city = tk.StringVar()


# Header
header = tk.Frame(root, bg=RED, height=78)
header.pack(fill="x")
header.pack_propagate(False)

tk.Label(
    header,
    text="BLOOD DONOR SYSTEM",
    font=("Arial", 20, "bold"),
    bg=RED,
    fg=WHITE
).pack(pady=(10, 0))

tk.Label(
    header,
    text="Donor Registration | Emergency Blood Search",
    font=("Arial", 10),
    bg=RED,
    fg=WHITE
).pack()


main = tk.Frame(root, bg=BG, padx=16, pady=10)
main.pack(fill="both", expand=True)


# Registration form
form = tk.LabelFrame(
    main,
    text=" Donor Registration ",
    font=("Arial", 12, "bold"),
    fg=DARK_RED,
    bg=WHITE,
    padx=12,
    pady=8
)
form.pack(fill="x")

for col in range(4):
    form.grid_columnconfigure(col, weight=1)


def make_label(text, row, col):
    tk.Label(
        form,
        text=text,
        font=("Arial", 10, "bold"),
        bg=WHITE,
        fg="black"
    ).grid(row=row, column=col, sticky="w", padx=6, pady=(5, 2))


def make_entry(var, row, col):
    tk.Entry(
        form,
        textvariable=var,
        font=("Arial", 10),
        relief="solid",
        bd=1
    ).grid(
        row=row, column=col,
        sticky="ew", padx=6, pady=(0, 5), ipady=4
    )


make_label("Full Name *", 0, 0)
make_entry(name_var, 1, 0)

make_label("Age (18-100) *", 0, 1)
make_entry(age_var, 1, 1)

make_label("Mobile Number *", 0, 2)
make_entry(mobile_var, 1, 2)

make_label("City *", 0, 3)
make_entry(city_var, 1, 3)

make_label("Blood Group *", 2, 0)

ttk.Combobox(
    form,
    textvariable=blood_var,
    values=GROUPS,
    state="readonly",
    font=("Arial", 10)
).grid(row=3, column=0, sticky="ew", padx=6, pady=3)

make_label("Last Donation (DD-MM-YYYY)", 2, 1)
make_entry(donation_var, 3, 1)

make_label("Available to Contact? *", 2, 2)

ttk.Combobox(
    form,
    textvariable=available_var,
    values=["Yes", "No"],
    state="readonly",
    font=("Arial", 10)
).grid(row=3, column=2, sticky="ew", padx=6, pady=3)

tk.Label(
    form,
    text="* Required fields",
    bg=WHITE,
    fg="gray",
    font=("Arial", 9)
).grid(row=3, column=3, sticky="w", padx=6)


# Registration buttons
buttons = tk.Frame(form, bg=WHITE)
buttons.grid(row=4, column=0, columnspan=4, pady=8)

tk.Button(
    buttons,
    text="Register Donor",
    command=register_donor,
    bg=RED,
    fg=WHITE,
    activebackground=DARK_RED,
    activeforeground=WHITE,
    font=("Arial", 10, "bold"),
    relief="flat",
    padx=15,
    pady=6,
    cursor="hand2"
).pack(side="left", padx=5)

tk.Button(
    buttons,
    text="Clear Form",
    command=clear_form,
    bg="lightgray",
    fg="black",
    font=("Arial", 10, "bold"),
    relief="flat",
    padx=15,
    pady=6,
    cursor="hand2"
).pack(side="left", padx=5)


# Emergency search section
search_box = tk.LabelFrame(
    main,
    text=" Emergency Blood Search ",
    font=("Arial", 12, "bold"),
    fg=DARK_RED,
    bg=WHITE,
    padx=10,
    pady=8
)
search_box.pack(fill="x", pady=9)

search_box.grid_columnconfigure(1, weight=1)
search_box.grid_columnconfigure(3, weight=2)

tk.Label(
    search_box,
    text="Blood Group",
    bg=WHITE,
    fg="black",
    font=("Arial", 10, "bold")
).grid(row=0, column=0, padx=5, pady=5)

ttk.Combobox(
    search_box,
    textvariable=search_blood,
    values=["All"] + GROUPS,
    state="readonly",
    font=("Arial", 10),
    width=12
).grid(row=0, column=1, padx=5, pady=5, sticky="ew")

tk.Label(
    search_box,
    text="City (optional)",
    bg=WHITE,
    fg="black",
    font=("Arial", 10, "bold")
).grid(row=0, column=2, padx=5, pady=5)

tk.Entry(
    search_box,
    textvariable=search_city,
    font=("Arial", 10),
    relief="solid",
    bd=1
).grid(
    row=0, column=3,
    padx=5, pady=5, sticky="ew", ipady=4
)


search_buttons = tk.Frame(search_box, bg=WHITE)
search_buttons.grid(row=1, column=0, columnspan=4, pady=5)

tk.Button(
    search_buttons,
    text="Search Available Donors",
    command=search_donors,
    bg=RED,
    fg=WHITE,
    font=("Arial", 10, "bold"),
    relief="flat",
    padx=12,
    pady=6,
    cursor="hand2"
).pack(side="left", padx=5)

tk.Button(
    search_buttons,
    text="Show All Donors",
    command=show_all,
    bg="pink",
    fg=DARK_RED,
    font=("Arial", 10, "bold"),
    relief="flat",
    padx=12,
    pady=6,
    cursor="hand2"
).pack(side="left", padx=5)


# Results table
result_box = tk.LabelFrame(
    main,
    text=" Donor Results ",
    font=("Arial", 12, "bold"),
    fg=DARK_RED,
    bg=WHITE,
    padx=6,
    pady=6
)
result_box.pack(fill="both", expand=True)

columns = ("name", "age", "city", "blood", "available")

table = ttk.Treeview(
    result_box,
    columns=columns,
    show="headings",
    height=7
)

for col, heading, width in [
    ("name", "Donor Name", 210),
    ("age", "Age", 65),
    ("city", "City", 170),
    ("blood", "Blood Group", 110),
    ("available", "Available", 110)
]:
    table.heading(col, text=heading)
    table.column(col, width=width, anchor="center")

scroll = ttk.Scrollbar(
    result_box,
    orient="vertical",
    command=table.yview
)

table.configure(yscrollcommand=scroll.set)
table.pack(side="left", fill="both", expand=True)
scroll.pack(side="right", fill="y")


# Status and footer
status = tk.Label(
    main,
    text="Register a donor or search for donors.",
    bg=BG,
    fg=DARK_RED,
    font=("Arial", 9, "bold"),
    anchor="w"
)
status.pack(fill="x", pady=(5, 0))

tk.Label(
    root,
    text="Demo directory only. Confirm donor consent, eligibility and availability.",
    bg=BG,
    fg="gray",
    font=("Arial", 8),
    pady=4
).pack(fill="x")

show_all()
root.mainloop()