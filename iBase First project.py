from openpyxl import Workbook, load_workbook
from datetime import datetime
from colorama import Fore, Style
import os
import random

file = "banking_system.xlsx"



def create_account():

    print(Fore.LIGHTMAGENTA_EX + "\n---------- CREATE ACCOUNT ----------" + Style.RESET_ALL)

    name = input("Enter Account Holder Name: ")
    mobile = input("Enter Mobile Number: ")
    pin = input("Enter 4-digit PIN: ")
    initial_deposit = float(input("Enter Initial Deposit: "))

    # Check if Excel file already exists
    if os.path.exists(file):

        wb = load_workbook(file)

        # Check if Account sheet exists
        if "Account" in wb.sheetnames:
            ws = wb["Account"]
        else:
            ws = wb.create_sheet("Account")

            ws.append([
                "Date",
                "Account Number",
                "Name",
                "Mobile Number",
                "PIN",
                "Balance"
            ])

    else:
        # Create Excel file for first account
        wb = Workbook()
        ws = wb.active
        ws.title = "Account"

        ws.append([
            "Date",
            "Account Number",
            "Name",
            "Mobile Number",
            "PIN",
            "Balance"
        ])

    # Generate unique account number
    while True:

        account_number = random.randint(100000, 999999)

        account_exists = False

        for row in range(2, ws.max_row + 1):

            if ws.cell(row, 2).value == account_number:
                account_exists = True
                break

        if not account_exists:
            break

    # Add new account
    ws.append([
        datetime.now().strftime("%d-%m-%Y %H:%M:%S"),
        account_number,
        name,
        mobile,
        pin,
        initial_deposit
    ])

    # Save Excel
    wb.save(file)
    wb.close()

    print(Fore.GREEN +
          "\nAccount created successfully!" +
          Style.RESET_ALL)

    print("Your Account Number is:", account_number)
    print("Your Balance is:", initial_deposit)



def deposite_money():
    print(Fore.LIGHTMAGENTA_EX + "----------- DEPOSITE MONEY ----------" + Style.RESET_ALL)
    
    ac = int(input("Enter Your Account Number : "))
    pinn = int(input("Enter Your PIN : "))

    # Noww open Excel file
    wb = load_workbook(file)
    ws = wb["Account"]

    # Search account
    found = False

    for row in range(2, ws.max_row + 1):

        account_number = ws.cell(row, 2).value
        pin = ws.cell(row, 5).value

        if account_number == ac and str(pin) == str(pinn):

            found = True

            amount = float(input("Enter Deposit Amount : "))

            if amount <= 0:
                print(Fore.RED + "Invalid Amount!" + Style.RESET_ALL)
                break

            old_balance = float(ws.cell(row, 6).value)

            new_balance = old_balance + amount

            # Update in Excel Sheet
            ws.cell(row, 6).value = new_balance

            wb.save(file)


            print(Fore.GREEN + "\nMoney Deposited Successfully!!" + Style.RESET_ALL)
            print("Deposited Amount : ", amount)
            print("Update Balance : ", new_balance)

            break

    if not found:
        print(Fore.RED + "\nInvalid Amount Number or PIN..!! Try Again..." + Style.RESET_ALL)

    wb.close()



def withdraw_money():
    print(Fore.LIGHTMAGENTA_EX + "---------- WITHDRAW MONEY ----------" + Style.RESET_ALL)

    ac = int(input("Enter Your Account Number : "))
    pinn = int(input("Enter Your PIN Number : "))

    # Open Excel File
    wb = load_workbook(file)
    ws = wb["Account"]

    found = False

    # Search account

    for row in range(2, ws.max_row + 1):

        account_number = ws.cell(row, 2).value
        pin = ws.cell(row, 5).value

        if account_number == ac and str(pin) == str(pinn):

            found = True

            amount = float(input("Enter Withdraw Amount : "))

            if amount <= 0:
                print(Fore.RED + "Invalid Amount!!" + Style.RESET_ALL)
                break

            # Get Current Balance
            current_balance = float(ws.cell(row, 6).value)

            # Check sufficient balance 
            if amount > current_balance:
                print(Fore.RED + "You don't have Sufficient Balance" + Style.RESET_ALL)
                print("Your Current Balance is : ", current_balance)
                break


            # Calculate new balance
            new_balance = current_balance - amount

            # Update balance in Excel
            ws.cell(row, 6).value = new_balance

            # save change
            wb.save(file)

            print(Fore.GREEN + "\nMoney withdrawn Successfully!!" + Style.RESET_ALL)
            print("Withdrawn Amount", amount)
            print("Remaining Balance", new_balance)

            break

    if not found:
        print(Fore.RED + "Invalid Account Number or PIN" + Style.RESET_ALL)
    wb.close()



def check_balance():
    print(Fore.LIGHTMAGENTA_EX + "---------- CHECK BALANCE ----------" + Style.RESET_ALL)

    ac = int(input("Enter Your Account Number : "))
    pinn = int(input("Enter Your PIN : "))

    # Open Excel File
    wb = load_workbook(file)
    ws = wb["Account"]

    found = False

    # Search account 
    for row in range(2, ws.max_row + 1):

        account_number = ws.cell(row, 2).value
        pin = ws.cell(row, 5).value

        if account_number == ac and str(pin) == str(pinn):

            found = True

            # Get account details
            name = ws.cell(row, 3).value
            balance = ws.cell(row, 6).value

            print(Fore.LIGHTBLACK_EX + "\n---------- ACCOUNT DETAILS ----------" + Style.RESET_ALL)
            print("Account Number : ", account_number)
            print("Account Holder : ", name)
            print("Current Balance : ", balance)

            break


    if not found:
        print(Fore.RED + "\nInvalid Account Number or PIN!")

    wb.close()




# BANK SYSTEM

print(Fore.CYAN + """
╔═════════════════════════════════════════════════╗
║     🏦 WELCOME TO BANK MANAGEMENT SYSTEM 🏦     ║
╚═════════════════════════════════════════════════╝
""")



def main():

    while True:
        print(Fore.LIGHTYELLOW_EX + """
        ╔══════════════════════════════════════╗
        ║              BANK MENU               ║
        ╠══════════════════════════════════════╣
        ║  1. 🗞️ Create Account                 ║
        ║  2. 💰 Deposit Money                 ║
        ║  3. 💸 Withdraw Money                ║
        ║  4. 💳 Check Balance                 ║
        ║  5. 🚪 Exit                          ║
        ╚══════════════════════════════════════╝
        """ + Style.RESET_ALL)
        
        choice = int(input("Enter your choice : "))

        if choice == 1:
            create_account()

        elif choice == 2:
            deposite_money()

        elif choice == 3:
            withdraw_money()

        elif choice == 4:
            check_balance()

        elif choice == 5:
            print(Fore.LIGHTWHITE_EX + "\nThank you for using our Banking System!" + Style.RESET_ALL)
            print("Have a Nice Day!")
            break

        else:
            print(Fore.RED + "Invalid Choice!! Try Again..." + Style.RESET_ALL)
            print("Please selecet 1 to 5.")


main()